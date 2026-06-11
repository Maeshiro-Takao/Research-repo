import pandas as pd
import statistics
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ====== パス設定 ======
csv_path = Path("/Users/maeshirotakao/ボートデータ/競走成績_グラフ/レース場別成績比較.csv")
output_path = Path("/Users/maeshirotakao/ボートデータ/競走成績_グラフ/レース場別成績比較.xlsx")

# ====== カラー定義 ======
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
ABOVE_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
BELOW_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

SEASONS = ["春(3-5月)", "夏(6-8月)", "秋(9-11月)", "冬(12-2月)"]

def load_data(csv_path):
    """CSVデータを読み込む"""
    print(f"読み込み中: {csv_path}")
    df = pd.read_csv(csv_path)
    
    avg_df = df[df['レース場'] == '全場平均'].copy()
    venue_df = df[df['レース場'] != '全場平均'].copy()
    venues = sorted(venue_df['レース場'].unique())
    
    return df, avg_df, venue_df, venues

def calculate_statistics(venue_df):
    """
    各場のデータをリストにして、statistics.pstdev()で母集団標準偏差を計算
    全場平均は使用しない
    """
    
    print("標準偏差を計算中（各場のデータから直接計算）...")
    
    statistics_dict = {}
    
    # ========== 全体の統計 ==========
    statistics_dict['全体'] = {}
    
    for boat in range(1, 7):
        statistics_dict['全体'][boat] = {}
        
        # 全場のこの艇の全体着率データをリストで取得
        all_venue_place = venue_df[(venue_df['タイプ'] == '全体着率') & (venue_df['艇'] == boat)]
        
        if len(all_venue_place) > 0:
            # 1着率のリストを作成して母集団標準偏差を計算
            values_1st = all_venue_place['1着率(%)'].values.tolist()
            mean_1st = statistics.mean(values_1st)
            std_1st = statistics.pstdev(values_1st, mu=mean_1st)
            statistics_dict['全体'][boat]['1着'] = {
                'mean': mean_1st,
                'std': std_1st,
                'values': values_1st
            }
            
            # 2着率のリストを作成して母集団標準偏差を計算
            values_2nd = all_venue_place['2着率(%)'].values.tolist()
            mean_2nd = statistics.mean(values_2nd)
            std_2nd = statistics.pstdev(values_2nd, mu=mean_2nd)
            statistics_dict['全体'][boat]['2着'] = {
                'mean': mean_2nd,
                'std': std_2nd,
                'values': values_2nd
            }
            
            # 3着率のリストを作成して母集団標準偏差を計算
            values_3rd = all_venue_place['3着率(%)'].values.tolist()
            mean_3rd = statistics.mean(values_3rd)
            std_3rd = statistics.pstdev(values_3rd, mu=mean_3rd)
            statistics_dict['全体'][boat]['3着'] = {
                'mean': mean_3rd,
                'std': std_3rd,
                'values': values_3rd
            }
    
    # 全場のこの艇の全体連帯率データをリストで取得
    for boat in range(1, 7):
        all_venue_rensai = venue_df[(venue_df['タイプ'] == '全体連帯率') & (venue_df['艇'] == boat)]
        
        if len(all_venue_rensai) > 0:
            # 2連帯率のリストを作成して母集団標準偏差を計算
            values_2ren = all_venue_rensai['2連帯率(%)'].values.tolist()
            mean_2ren = statistics.mean(values_2ren)
            std_2ren = statistics.pstdev(values_2ren, mu=mean_2ren)
            statistics_dict['全体'][boat]['2連帯'] = {
                'mean': mean_2ren,
                'std': std_2ren,
                'values': values_2ren
            }
            
            # 3連帯率のリストを作成して母集団標準偏差を計算
            values_3ren = all_venue_rensai['3連帯率(%)'].values.tolist()
            mean_3ren = statistics.mean(values_3ren)
            std_3ren = statistics.pstdev(values_3ren, mu=mean_3ren)
            statistics_dict['全体'][boat]['3連帯'] = {
                'mean': mean_3ren,
                'std': std_3ren,
                'values': values_3ren
            }
    
    # ========== 季節別の統計 ==========
    for season in SEASONS:
        statistics_dict[season] = {}
        
        for boat in range(1, 7):
            statistics_dict[season][boat] = {}
            
            # 全場のこの艇の季節別着率データをリストで取得
            all_venue_place = venue_df[(venue_df['タイプ'] == '季節別着率') & 
                                       (venue_df['季節'] == season) & 
                                       (venue_df['艇'] == boat)]
            
            if len(all_venue_place) > 0:
                # 1着率のリストを作成して母集団標準偏差を計算
                values_1st = all_venue_place['1着率(%)'].values.tolist()
                mean_1st = statistics.mean(values_1st)
                std_1st = statistics.pstdev(values_1st, mu=mean_1st)
                statistics_dict[season][boat]['1着'] = {
                    'mean': mean_1st,
                    'std': std_1st,
                    'values': values_1st
                }
                
                # 2着率のリストを作成して母集団標準偏差を計算
                values_2nd = all_venue_place['2着率(%)'].values.tolist()
                mean_2nd = statistics.mean(values_2nd)
                std_2nd = statistics.pstdev(values_2nd, mu=mean_2nd)
                statistics_dict[season][boat]['2着'] = {
                    'mean': mean_2nd,
                    'std': std_2nd,
                    'values': values_2nd
                }
                
                # 3着率のリストを作成して母集団標準偏差を計算
                values_3rd = all_venue_place['3着率(%)'].values.tolist()
                mean_3rd = statistics.mean(values_3rd)
                std_3rd = statistics.pstdev(values_3rd, mu=mean_3rd)
                statistics_dict[season][boat]['3着'] = {
                    'mean': mean_3rd,
                    'std': std_3rd,
                    'values': values_3rd
                }
        
        # 全場のこの艇の季節別連帯率データをリストで取得
        for boat in range(1, 7):
            all_venue_rensai = venue_df[(venue_df['タイプ'] == '季節別連帯率') & 
                                        (venue_df['季節'] == season) & 
                                        (venue_df['艇'] == boat)]
            
            if len(all_venue_rensai) > 0:
                # 2連帯率のリストを作成して母集団標準偏差を計算
                values_2ren = all_venue_rensai['2連帯率(%)'].values.tolist()
                mean_2ren = statistics.mean(values_2ren)
                std_2ren = statistics.pstdev(values_2ren, mu=mean_2ren)
                statistics_dict[season][boat]['2連帯'] = {
                    'mean': mean_2ren,
                    'std': std_2ren,
                    'values': values_2ren
                }
                
                # 3連帯率のリストを作成して母集団標準偏差を計算
                values_3ren = all_venue_rensai['3連帯率(%)'].values.tolist()
                mean_3ren = statistics.mean(values_3ren)
                std_3ren = statistics.pstdev(values_3ren, mu=mean_3ren)
                statistics_dict[season][boat]['3連帯'] = {
                    'mean': mean_3ren,
                    'std': std_3ren,
                    'values': values_3ren
                }
    
    return statistics_dict

def calculate_venue_deviations(venue, venue_df, statistics_dict, season_type='全体'):
    """
    指定された場の偏差を計算
    偏差 = 値 - 全場平均（単純な差）
    """
    
    deviations = {}
    
    for boat in range(1, 7):
        deviations[boat] = {
            '1着': None,
            '2着': None,
            '3着': None,
            '2連帯': None,
            '3連帯': None,
        }
        
        # ========== 着率の偏差を計算 ==========
        if season_type == '全体':
            place_venue = venue_df[(venue_df['レース場'] == venue) & 
                                   (venue_df['タイプ'] == '全体着率') & 
                                   (venue_df['艇'] == boat)]
        else:
            place_venue = venue_df[(venue_df['レース場'] == venue) & 
                                   (venue_df['タイプ'] == '季節別着率') & 
                                   (venue_df['季節'] == season_type) &
                                   (venue_df['艇'] == boat)]
        
        if not place_venue.empty:
            # 1着率の偏差
            value_1st = place_venue.iloc[0]['1着率(%)']
            mean_1st = statistics_dict[season_type][boat]['1着']['mean']
            deviations[boat]['1着'] = value_1st - mean_1st
            
            # 2着率の偏差
            value_2nd = place_venue.iloc[0]['2着率(%)']
            mean_2nd = statistics_dict[season_type][boat]['2着']['mean']
            deviations[boat]['2着'] = value_2nd - mean_2nd
            
            # 3着率の偏差
            value_3rd = place_venue.iloc[0]['3着率(%)']
            mean_3rd = statistics_dict[season_type][boat]['3着']['mean']
            deviations[boat]['3着'] = value_3rd - mean_3rd
        
        # ========== 連帯率の偏差を計算 ==========
        if season_type == '全体':
            rensai_venue = venue_df[(venue_df['レース場'] == venue) & 
                                    (venue_df['タイプ'] == '全体連帯率') & 
                                    (venue_df['艇'] == boat)]
        else:
            rensai_venue = venue_df[(venue_df['レース場'] == venue) & 
                                    (venue_df['タイプ'] == '季節別連帯率') & 
                                    (venue_df['季節'] == season_type) &
                                    (venue_df['艇'] == boat)]
        
        if not rensai_venue.empty:
            # 2連帯率の偏差
            value_2ren = rensai_venue.iloc[0]['2連帯率(%)']
            mean_2ren = statistics_dict[season_type][boat]['2連帯']['mean']
            deviations[boat]['2連帯'] = value_2ren - mean_2ren
            
            # 3連帯率の偏差
            value_3ren = rensai_venue.iloc[0]['3連帯率(%)']
            mean_3ren = statistics_dict[season_type][boat]['3連帯']['mean']
            deviations[boat]['3連帯'] = value_3ren - mean_3ren
    
    return deviations

def add_deviation_sheet(wb, venue, venue_df, statistics_dict, season_type='全体'):
    """偏差シートを追加"""
    
    deviations = calculate_venue_deviations(venue, venue_df, statistics_dict, season_type)
    
    # シート名を決定
    if season_type == '全体':
        sheet_name = venue
    else:
        season_short = season_type.split('(')[0]
        sheet_name = f"{venue}_{season_short}"
    
    if len(sheet_name) > 31:
        sheet_name = sheet_name[:31]
    
    ws = wb.create_sheet(sheet_name)
    
    headers = ['艇', '1着偏差', '2着偏差', '3着偏差', '2連偏差', '3連偏差']
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = header
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER
    
    row = 2
    for boat in range(1, 7):
        cell = ws.cell(row=row, column=1)
        cell.value = boat
        cell.border = BORDER
        cell.alignment = Alignment(horizontal='center')
        
        dev = deviations[boat]
        
        col_data = [
            ('1着', 2),
            ('2着', 3),
            ('3着', 4),
            ('2連帯', 5),
            ('3連帯', 6),
        ]
        
        for dev_key, col_idx in col_data:
            cell = ws.cell(row=row, column=col_idx)
            
            if dev[dev_key] is not None:
                cell.value = dev[dev_key]
                cell.border = BORDER
                cell.alignment = Alignment(horizontal='center')
                cell.number_format = '0.00'
                
                if dev[dev_key] > 0:
                    cell.fill = ABOVE_FILL
                elif dev[dev_key] < 0:
                    cell.fill = BELOW_FILL
            else:
                cell.border = BORDER
        
        row += 1
    
    ws.column_dimensions['A'].width = 8
    for col in range(2, 7):
        ws.column_dimensions[get_column_letter(col)].width = 14

def create_comparison_excel(csv_path, output_path):
    """Excelファイルを生成"""
    
    df, avg_df, venue_df, venues = load_data(csv_path)
    statistics_dict = calculate_statistics(venue_df)
    
    print("Excelファイルを生成中...")
    
    wb = Workbook()
    wb.remove(wb.active)
    
    # ============================================
    # [1] 平均値シート
    # ============================================
    print("  - 平均値シート")
    ws_avg = wb.create_sheet("平均値")
    
    headers = ['艇', '1着率(%)', '2着率(%)', '3着率(%)', '2連帯率(%)', '3連帯率(%)']
    for col_idx, header in enumerate(headers, 1):
        cell = ws_avg.cell(row=1, column=col_idx)
        cell.value = header
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER
    
    row = 2
    for boat in range(1, 7):
        cell = ws_avg.cell(row=row, column=1)
        cell.value = boat
        cell.border = BORDER
        cell.alignment = Alignment(horizontal='center')
        
        # 統計情報から平均値を取得
        for col, metric, metric_name in [(2, '1着', '1着率(%)'), (3, '2着', '2着率(%)'), (4, '3着', '3着率(%)')]:
            cell = ws_avg.cell(row=row, column=col)
            cell.value = statistics_dict['全体'][boat][metric]['mean']
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='center')
            cell.number_format = '0.00'
        
        for col, metric, metric_name in [(5, '2連帯', '2連帯率(%)'), (6, '3連帯', '3連帯率(%)')]:
            cell = ws_avg.cell(row=row, column=col)
            cell.value = statistics_dict['全体'][boat][metric]['mean']
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='center')
            cell.number_format = '0.00'
        
        row += 1
    
    ws_avg.column_dimensions['A'].width = 8
    for col in range(2, 7):
        ws_avg.column_dimensions[get_column_letter(col)].width = 14
    
    # ============================================
    # [2] 統計情報シート（全体）
    # ============================================
    print("  - 統計情報シート（全体）")
    ws_stats = wb.create_sheet("統計情報_全体")
    
    headers = ['艇', '1着_平均', '1着_標準偏差', '2着_平均', '2着_標準偏差', 
               '3着_平均', '3着_標準偏差', '2連_平均', '2連_標準偏差', 
               '3連_平均', '3連_標準偏差']
    for col_idx, header in enumerate(headers, 1):
        cell = ws_stats.cell(row=1, column=col_idx)
        cell.value = header
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER
    
    row = 2
    for boat in range(1, 7):
        cell = ws_stats.cell(row=row, column=1)
        cell.value = boat
        cell.border = BORDER
        cell.alignment = Alignment(horizontal='center')
        
        col_idx = 2
        for metric in ['1着', '2着', '3着', '2連帯', '3連帯']:
            cell = ws_stats.cell(row=row, column=col_idx)
            cell.value = statistics_dict['全体'][boat][metric]['mean']
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='center')
            cell.number_format = '0.00'
            col_idx += 1
            
            cell = ws_stats.cell(row=row, column=col_idx)
            cell.value = statistics_dict['全体'][boat][metric]['std']
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='center')
            cell.number_format = '0.00'
            col_idx += 1
        
        row += 1
    
    ws_stats.column_dimensions['A'].width = 8
    for col in range(2, 12):
        ws_stats.column_dimensions[get_column_letter(col)].width = 12
    ws_stats.row_dimensions[1].height = 25
    
    # ============================================
    # [2-2] 統計情報シート（季節別）
    # ============================================
    print("  - 統計情報シート（季節別）")
    
    for season in SEASONS:
        season_short = season.split('(')[0]
        ws_season = wb.create_sheet(f"統計情報_{season_short}")
        
        headers = ['艇', '1着_平均', '1着_標準偏差', '2着_平均', '2着_標準偏差', 
                   '3着_平均', '3着_標準偏差', '2連_平均', '2連_標準偏差', 
                   '3連_平均', '3連_標準偏差']
        for col_idx, header in enumerate(headers, 1):
            cell = ws_season.cell(row=1, column=col_idx)
            cell.value = header
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = BORDER
        
        row = 2
        for boat in range(1, 7):
            cell = ws_season.cell(row=row, column=1)
            cell.value = boat
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='center')
            
            col_idx = 2
            for metric in ['1着', '2着', '3着', '2連帯', '3連帯']:
                cell = ws_season.cell(row=row, column=col_idx)
                cell.value = statistics_dict[season][boat][metric]['mean']
                cell.border = BORDER
                cell.alignment = Alignment(horizontal='center')
                cell.number_format = '0.00'
                col_idx += 1
                
                cell = ws_season.cell(row=row, column=col_idx)
                cell.value = statistics_dict[season][boat][metric]['std']
                cell.border = BORDER
                cell.alignment = Alignment(horizontal='center')
                cell.number_format = '0.00'
                col_idx += 1
            
            row += 1
        
        ws_season.column_dimensions['A'].width = 8
        for col in range(2, 12):
            ws_season.column_dimensions[get_column_letter(col)].width = 12
        ws_season.row_dimensions[1].height = 25
    
    # ============================================
    # [3] 各レース場別シート（全体と季節別）
    # ============================================
    print("  - 各レース場別シート（全体+季節別）")
    
    for venue in venues:
        add_deviation_sheet(wb, venue, venue_df, statistics_dict, '全体')
        for season in SEASONS:
            add_deviation_sheet(wb, venue, venue_df, statistics_dict, season)
    
    wb.save(output_path)
    print(f"\n✅ Excel生成完了: {output_path}")

if __name__ == "__main__":
   
    try:
        create_comparison_excel(csv_path, output_path)
        print("\n" + "="*70)
        print("\n計算方法:")
        print("  1. 各場の着率データをリスト化")
        print("  2. statistics.mean()で平均値を計算")
        print("  3. statistics.pstdev()で母集団標準偏差を計算")
        print("  4. 偏差 = レース場の値 - 全場の平均値")
        print("="*70)
    except Exception as e:
        print(f"❌ エラー: {e}")
        import traceback
        traceback.print_exc()